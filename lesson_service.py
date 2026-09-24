import os
import re
import json
import asyncio
from typing import Any, List, Dict, Optional, Tuple
from google import genai
from google.genai import types
from dotenv import load_dotenv

import pdfplumber
from mutagen.mp3 import MP3

from chat_service import (
    save_or_update_lesson,
    get_lesson,
    list_lessons,
    delete_lesson,
    sanitize_project_name,
    get_chroma_client,
    get_ai_client,
)

load_dotenv()


# --- MULTIMODÁLNÍ INGESCE A ZPRACOVÁNÍ ---

def detect_mime_type(filename: str) -> str:
    ext = filename.lower().split(".")[-1]
    mime_map = {
        "pdf": "application/pdf",
        "png": "image/png",
        "jpg": "image/jpeg",
        "jpeg": "image/jpeg",
        "webp": "image/webp",
        "gif": "image/gif",
        "txt": "text/plain",
        "md": "text/markdown",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    }
    return mime_map.get(ext, "application/octet-stream")


def normalize_image_part_bytes(blob: bytes, original_mime: str) -> Optional[Tuple[bytes, str]]:
    """
    Zajistí, že obrazová data mají validní MIME type podporovaný Google Gemini API
    (image/jpeg, image/png, image/webp, image/heic, image/heif). Pokud je formát
    nekompatibilní (TIFF, BMP, GIF apod.), pokusí se jej zkonvertovat přes PIL na JPEG/PNG.
    Nepodporované formáty (např. vektorový WMF / EMF ze slidů), které nelze převést, bezpečně přeskočí,
    aby nezpůsobily chybu 400 INVALID_ARGUMENT.
    """
    supported = {"image/jpeg", "image/png", "image/webp", "image/heic", "image/heif"}
    norm_mime = (original_mime or "").lower().strip()

    # Ignorujeme vektorové Windows metafiles (WMF / EMF), které Gemini odmítá s HTTP 400
    if "wmf" in norm_mime or "emf" in norm_mime:
        return None

    if norm_mime in supported:
        return blob, norm_mime

    # Pokusíme se převést TIFF, BMP, GIF apod. na JPEG/PNG přes PIL
    try:
        from PIL import Image
        import io
        img = Image.open(io.BytesIO(blob))
        out = io.BytesIO()
        if img.mode in ("RGBA", "LA", "P"):
            img.convert("RGBA").save(out, format="PNG")
            return out.getvalue(), "image/png"
        else:
            img.convert("RGB").save(out, format="JPEG", quality=88)
            return out.getvalue(), "image/jpeg"
    except Exception:
        return None


async def extract_multimodal_parts(
    file_paths: List[Tuple[str, str]]
) -> Tuple[List[Any], str, List[Dict[str, Any]]]:
    """
    Zpracuje nahrané soubory do seznamu Gemini Part objektů a extrahovaného textu.
    file_paths je seznam dvojic (cesta_k_souboru, puvodni_nazev).
    """
    parts = []
    combined_text = ""
    file_metadata = []

    for file_path, original_filename in file_paths:
        mime = detect_mime_type(original_filename)
        ext = original_filename.lower().split(".")[-1]

        with open(file_path, "rb") as f:
            file_bytes = f.read()

        file_metadata.append(
            {
                "filename": original_filename,
                "mime_type": mime,
                "size_bytes": len(file_bytes),
            }
        )

        # PDF předáváme přímo nativně do Gemini přes Part.from_bytes
        if mime == "application/pdf":
            try:
                part = types.Part.from_bytes(data=file_bytes, mime_type=mime)
                parts.append(part)
            except Exception as e:
                print(f"[WARN] Nelze vytvořit Part.from_bytes pro {original_filename}: {e}")

        # Samostatné obrázky normalizujeme pro Gemini API
        elif mime.startswith("image/"):
            norm = normalize_image_part_bytes(file_bytes, mime)
            if norm:
                norm_bytes, norm_mime = norm
                try:
                    part = types.Part.from_bytes(data=norm_bytes, mime_type=norm_mime)
                    parts.append(part)
                except Exception as e:
                    print(f"[WARN] Nelze vytvořit Part.from_bytes pro obrázek {original_filename}: {e}")

        # Pro textové formáty a PDF se pokusíme vytáhnout i text pro RAG indexaci
        if ext == "pdf":
            try:
                with pdfplumber.open(file_path) as pdf:
                    for idx, page in enumerate(pdf.pages):
                        text = page.extract_text()
                        if text:
                            combined_text += f"\n--- [{original_filename} str. {idx + 1}] ---\n{text}\n"
            except Exception as e:
                print(f"[WARN] pdfplumber selhal pro {original_filename}: {e}")

        elif ext == "txt" or ext == "md":
            try:
                text = file_bytes.decode("utf-8", errors="ignore")
                combined_text += f"\n--- [{original_filename}] ---\n{text}\n"
            except Exception as e:
                print(f"[WARN] TXT čtení selhalo: {e}")

        elif ext == "docx":
            try:
                import docx
                doc = docx.Document(file_path)
                docx_text = "\n".join([p.text for p in doc.paragraphs if p.text])
                combined_text += f"\n--- [{original_filename}] ---\n{docx_text}\n"
            except Exception as e:
                print(f"[WARN] DOCX čtení selhalo: {e}")

        elif ext == "pptx":
            try:
                import pptx
                prs = pptx.Presentation(file_path)
                pptx_text = ""
                for s_idx, slide in enumerate(prs.slides):
                    slide_text = []
                    for shape in slide.shapes:
                        if hasattr(shape, "text") and shape.text:
                            slide_text.append(shape.text)
                    if slide_text:
                        pptx_text += f"\n[Slide {s_idx + 1}]: " + " | ".join(slide_text)
                combined_text += f"\n--- [{original_filename}] ---\n{pptx_text}\n"

                # Extrakce obrázků a grafů ze slidů (např. fotografie otázek, kazuistik, diagramy)
                pptx_images_added = 0
                for s_idx, slide in enumerate(prs.slides):
                    for shape in slide.shapes:
                        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
                            try:
                                img = shape.image
                                blob = img.blob
                                # Filtrujeme drobné dekorace/ikony (< 10 KB) a omezíme celkový počet částí
                                if len(blob) >= 10240 and len(parts) < 35:
                                    norm = normalize_image_part_bytes(blob, img.content_type or "image/png")
                                    if norm:
                                        norm_bytes, norm_mime = norm
                                        part = types.Part.from_bytes(data=norm_bytes, mime_type=norm_mime)
                                        parts.append(part)
                                        pptx_images_added += 1
                            except Exception:
                                pass
                if pptx_images_added > 0:
                    print(f"[INFO] Přidáno {pptx_images_added} validních obrázků ze slidů prezentace '{original_filename}'.")
            except Exception as e:
                print(f"[WARN] PPTX čtení selhalo: {e}")

    return parts, combined_text, file_metadata


async def index_lesson_text_to_chroma(
    project_id: str, combined_text: str, lesson_title: str
):
    """Zaindexuje extrahovaný text a osnovu do ChromaDB kolekce projektu."""
    if not combined_text.strip():
        return

    from main import chunk_text

    chroma = get_chroma_client()
    collection_name = f"proj_{sanitize_project_name(project_id)}"
    collection = chroma.get_or_create_collection(name=collection_name)

    chunks = chunk_text(combined_text)
    ids = [f"lesson_{sanitize_project_name(lesson_title)}_chunk_{i}" for i in range(len(chunks))]
    metadatas = [{"source": f"Lekce: {lesson_title}"} for _ in range(len(chunks))]

    collection.add(documents=chunks, metadatas=metadatas, ids=ids)
    try:
        from chat_service import index_chunks_to_fts
        fts_chunks = [
            {
                "id": ids[i],
                "document": chunks[i],
                "source": metadatas[i]["source"],
                "page": "1",
            }
            for i in range(len(chunks))
        ]
        index_chunks_to_fts(project_id, fts_chunks)
    except Exception as e:
        print(f"[lesson_service] FTS lesson indexing warning: {e}")


# --- DVOUFÁZOVÁ PEDAGOGICKÁ PIPELINE ---

STAGE_A_STUDY_MATERIAL_PROMPT = """Jsi špičkový profesor medicíny a excelentní akademický pedagog.
Tvým úkolem je na základě přiložených multimediálních materiálů (které mohou být v cizím jazyce, např. francouzštině, angličtině, němčině) vytvořit komplexní, vysoce strukturovaný, rigorózní a pedagogicky dokonalý studijní text k tématu: "{TOPIC}".

CÍLOVÝ JAZYK STUDIJNÍHO TEXTU: {TARGET_LANGUAGE} (veškerý text musí být v tomto jazyce se správnou odbornou terminologií).

POVINNÁ STRUKTURA A METODIKA (MARKDOWN):
Text musí být podrobný, jasně členěný, scannovatelný a nabitý informacemi (žádné prázdné obecnosti, žádný zdlouhavý úvod, žádný "wall of text"). Použij tuto přesnou osnovu:

# {TOPIC}

## 1. Konceptuální rámec a definice
- BEZ ZDLOUHAVÉHO ÚVODU: Žádné obecné tlachání ani zbytečný úvodní odstavec. Začni okamžitě definicí a meritem věci.
- Přesná odborná definice
- Epidemiologický a klinický kontext, zařazení nozologické jednotky

## 2. Patofyziologie a biologické mechanismy
- Detailní logický řetězec vzniku a progrese
- Buněčné, humorální a systémové mechanismy (včetně popisu diagramů a schémat z materiálů)
- Následky na cílové orgány

## 3. Klinický obraz a symptomatologie
- Typické symptomy a časné příznaky
- Pozdní a atypické manifestace
- Fyzikální nález

## 4. Diagnostický algoritmus a klasifikace
- Iniciální laboratorní a zobrazovací vyšetření
- Zlatý diagnostický standard
- SROVNÁVACÍ MARKDOWN TABULKA: Vytvoř přehlednou tabulku pro klasifikaci stádií nebo diferenciální diagnostiku (sloupce: Jednotka / Stádium | Klíčové znaky | Diagnostický nález | Typická úskalí).

## 5. Terapeutický management a farmakoterapie
- Nežádoucí vlivy a režimová opatření
- Farmakoterapie: Léky první volby, mechanismus účinku, konkrétní zástupci a kontraindikace
- Intervenční, chirurgické nebo akutní postupy

## 6. Klinické perly, chytáky a Red Flags
- 3–5 zásadních klinických pastí a varovných příznaků (Red Flags)
- Nejčastější chyby u zkoušek a v klinické praxi

## 7. Rychlý High-Yield souhrn
- Heslovitý přehled 5–7 nejdůležitějších bodů pro rychlé zopakování

## 8. Testové otázky z výchozích materiálů
⚠️ KRITICKÉ PRAVIDLO O EXISTENCI OTÁZEK:
- Tuto sekci 8 vytvoř POUZE A JENOM TEHDY, pokud nahrané podklady (PDF, prezentace, texty, fotky slidů) SKUTEČNĚ OBSAHUJÍ testové otázky (otázky s volbou odpovědí MCQ s možnostmi A/B/C/D/E, případně otevřené kontrolní/zkouškové otázky či kazuistiky)!
- POKUD VE VÝCHOZÍCH DOKUMENTECH ŽÁDNÉ TESTOVÉ OTÁZKY NEJSOU:
  * Sekci 8 VŮBEC NEVYTVÁŘEJ!
  * Text studijní lekce končí sekcí 7 (Rychlý High-Yield souhrn).
  * Za žádných okolností si nevymýšlej žádné vlastní umělé testové otázky!

DŮRAZNÉ PRAVIDLO PRO FOTOGRAFIE SLIDŮ Z PŘEDNÁŠEK A ZAKROUŽKOVANÉ ODPOVĚDI:
- Podklady velmi často obsahují FOTOGRAFIE PROMÍTANÝCH SLIDŮ Z PŘEDNÁŠEK A SEMINÁŘŮ (např. kazuistiky „Cas clinique“, kontrolní otázky a testy).
- Na těchto snímcích bývají správné odpovědi VYZNAČENY RUČNĚ (např. perem/fixem/tužkou zakroužkované písmeno možnosti jako (B) nebo (C), podtržení volby, zaškrtnutí či rukopisná poznámka).
- PROZKOUMEJ VŠECHNY STRÁNKY A SNÍMKY DOKUMENTU AŽ DO ÚPLNÉHO KONCE (testové otázky a kazuistiky bývají typicky na závěrečných snímcích prezentace)!
- Pokud na fotkách slidů najdeš jakékoliv otázky a zakroužkované/označené odpovědi, MUSÍŠ JE VŠECHNY PŘEVZÍT DO TÉTO SEKCE 8!

POKUD JSOU V PODKLADECH TESTOVÉ OTÁZKY PŘÍTOMNY:
Přidej je nakonec lekce přesně tak, jak jsou v originále (zachovej původní znění otázek, pořadí a varianty odpovědí; pokud jsou cizojazyčné, např. francouzsky, přelož je věrně do {TARGET_LANGUAGE} se zachováním přesného smyslu a odborných termínů).
Struktura musí umožnit uživateli buď si test samostatně vyplnit, nebo rovnou nahlédnout na správné odpovědi a vysvětlení rozkliknutím interaktivního bloku `<details>`.

ZPRACOVÁNÍ ODPOVĚDÍ A VYSVĚTLENÍ:
1. Pokud výchozí text či fotka slidu OBSAHUJE označené / zakroužkované odpovědi:
   - Uveď tuto zakroužkovanou možnost jako správnou odpověď.
   - PŘIDEJ PODROBNÉ VYSVĚTLENÍ KAŽDÉ MOŽNOSTI:
     * Proč je zakroužkovaná volba správná (jaký biologický/klinický fakt či mechanismus ji potvrzuje). Pokud je na slidu rukopisná poznámka (např. „bilan avant transfusion“, „flux portal“), zapracuj ji do vysvětlení!
     * Proč je každá ze špatných voleb nesprávná (v čem spočívá omyl, chyták či záměna).
   - Označ původ: `*(Zdroj odpovědi: Zakroužkováno / vyznačeno přímo v originálních podkladech)*`.

2. Pokud výchozí materiál obsahuje otázky, ale NEJSOU v něm vyznačeny odpovědi (žádný kroužek ani klíč):
   - Jako excelentní pedagog a lékař správnou odpověď i vysvětlení dovygeneruj.
   - PŘIDEJ PODROBNÉ VYSVĚTLENÍ KAŽDÉ MOŽNOSTI (proč je správná volba správně a proč jsou ostatní možnosti špatně).
   - JASNĚ A ZŘETELNĚ OZNAČ, ŽE JDE O ODPOVĚĎ VYGENEROVANOU AI: `*(⚠️ Odpověď a řešení vygenerované AI – výchozí dokument neobsahoval vyznačené odpovědi)*`.

PŘESNÉ FORMÁTOVÁNÍ PRO OTÁZKY S VÝBĚREM MOŽNOSTÍ (MCQ):
### Otázka 1: [Přesné znění otázky z podkladů]
- **A)** [Znění možnosti A]
- **B)** [Znění možnosti B]
- **C)** [Znění možnosti C]
- **D)** [Znění možnosti D]

<details>
<summary>🔍 Zobrazit správnou odpověď a vysvětlení</summary>

**Správná odpověď:** [Např. B]
*(Zdroj odpovědi: Klíč v originálním podkladu / NEBO: ⚠️ Odpověď vygenerovaná AI – výchozí dokument neobsahoval klíč)*

**Podrobné vysvětlení jednotlivých možností:**
- **A (Špatně):** [Proč je možnost A nesprávná]
- **B (Správně):** [Proč je možnost B správná a na jakém mechanismu staví]
- **C (Špatně):** [Proč je možnost C nesprávná]
- **D (Špatně):** [Proč je možnost D nesprávná]
</details>

PŘESNÉ FORMÁTOVÁNÍ PRO OTEVŘENÉ KONTROLNÍ OTÁZKY:
### Otázka X: [Přesné znění otázky z podkladů]

<details>
<summary>🔍 Zobrazit vzorovou odpověď a vysvětlení</summary>

**Vzorová odpověď a řešení:**
*(Zdroj odpovědi: Originální podklad / NEBO: ⚠️ Vzorová odpověď vygenerovaná AI – výchozí dokument neobsahoval klíč)*

[Detailní odborné řešení, klíčové argumenty, diferenciální diagnostika či terapeutický postup, který musí odpověď obsahovat.]
</details>

STRIKTNÍ POKYNY PRO FORMÁTOVÁNÍ:
- Používej tučné písmo pro klíčové termíny a léky.
- Zachovej přesnost všech čísel, skórovacích schémat a kritérií z materiálů.
- Nezačínej žádným úvodním povídáním jako "Zde je váš text", začni přímo hlavním nadpisem `# {TOPIC}`."""


STAGE_B_LECTURE_SCRIPT_PROMPT = """Jsi charismatický univerzitní profesor a seniorní klinický pedagog s pověstí nejlepšího a nejpoutavějšího přednášejícího na lékařské fakultě.
Před tebou leží detailní písemný studijní text k tématu: "{TOPIC}". Student má tento text otevřený před sebou na obrazovce.

Tvým úkolem je vytvořit samostatný MLUVENÝ PŘEDNÁŠKOVÝ SCÉNÁŘ (Spoken Lecture Script) v jazyce: {TARGET_LANGUAGE}.

KRITICKY DŮLEŽITÉ ZÁSADY PŘEDNÁŠKY:

1. VYNECH ZDLOUHAVÝ ÚVOD (PŘÍMÝ A ÚDERNÝ START):
   - ŽÁDNÉ zdlouhavé uvítací řeči, formální ceremonie, představování sylabu ani prázdná omáčka typu: "Vítám vás na dnešní přednášce, je mi velkým potěšením, že se dnes scházíme nad tímto fascinujícím tématem, dnes se podíváme na...".
   - Úvod omez na maximálně JEDNU jedinou svižnou větu (např. "Dobrý den, pojďme přímo k jádru věci – k {TOPIC}." nebo "Dobrý den, otevřete si studijní text, začínáme rovnou klinickým jádrem.") a OKAMŽITĚ jdi do výkladu a podstaty klinického problému!

2. AUTENTICKÁ VÝUKOVÁ PŘEDNÁŠKA (NIKOLIV AUDIOKNIHA):
   - Tato nahrávka NESMÍ být doslovným čtením studijního textu!
   - Posluchač nepotřebuje mechanické předčítání odrážek a tabulek. Potřebuje špičkovou vysokoškolskou přednášku z auly – živou, energickou, srozumitelnou a plnou hlubokých klinických souvislostí ("PROČ věci fungují tak, jak fungují").
   - Mluv energickým tónem zkušeného klinika, pokládej řečnické otázky ("Proč se to vlastně děje?", "Co v tu chvíli uděláte na příjmu?"), dávej látce přirozenou dynamiku a pedagogický tah na branku.

3. AKTIVNÍ INTERAKTIVNÍ ODKAZY NA STUDIJNÍ TEXT:
   - Průběžně studenta přirozeně naváděj na pasáže v textu, který má otevřený před sebou:
     (např. "Když se teď podíváte do sekce 2 na to patofyziologické schéma...",
     "Všimněte si ve čtvrté kapitole v naší srovnávací tabulce třetího sloupce – to je ten zásadní zlom...",
     "V sekci farmakoterapie máte vypsané přesné zástupce léků, ale já vám chci předat tu základní intuici pro jejich volbu v praxi...").

4. INTUITIVNÍ ANALOGIE A METAFORY:
   - Nejsložitější biologické a patofyziologické děje osvětli pomocí trefné analogie ze života (např. dopravní kolaps, přetlakový bezpečnostní ventil, ucpané potrubí, elektrický zkrat apod.).

5. KLINICKÁ REALITA, RED FLAGS A ZKOUŠKOVÉ CHYTÁKY:
   - Zdůrazni fatální pasti a omyly mediků u zkoušek a mladých lékařů na pohotovosti ("Na tohle si dejte obrovský pozor...", "Tady u zkoušky většina lidí zaváhá...").

6. ODKAZ NA TESTOVÉ OTÁZKY NA KONCI (NEPŘEDČÍTAT):
   - Pokud studijní text obsahuje na konci testové otázky z původních materiálů, v závěrečné kapitole přednášky studenta pouze stručně vyzvi, aby si je prošel a zkontroloval rozbory odpovědí ("Na samém konci studijního textu máte k dispozici původní testové otázky z vašich podkladů s podrobným vysvětlením každé možnosti – určitě si je po přednášce zkuste sami projít a ověřit si své znalosti.").
   - V audio přednášce NIKDY mechanicky nečti celé otázky ani možnosti A/B/C/D.

7. JASNÉ ČLENĚNÍ DO KAPITOL (PRO AUDIO STOPU):
   - Přednášku striktně rozděl do 5 až 6 logických kapitol.
   - Každá kapitola MUSÍ začínat PŘESNĚ tímto značkovačem na samostatném řádku:
     [KAPITOLA: 01 | Klinický kontext a jádro problému]
     [KAPITOLA: 02 | Klíčová patofyziologie a intuitivní analogie]
     [KAPITOLA: 03 | Rozbor diagnostiky a pohled do tabulek]
     [KAPITOLA: 04 | Terapeutická strategie a klinické perly]
     [KAPITOLA: 05 | Závěrečné shrnutí a doporučení do praxe]

8. MLUVNÍ KULTURA A DÉLKA:
   - Piš v mluveném, energickém, kultivovaném a pedagogicky poutavém jazyce.
   - Délka textu by měla odpovídat souvislému cca 8–15minutovému výkladu (cca 1 200 až 2 500 slov)."""


def extract_chapters_from_script(
    script: str, total_duration_seconds: float = 600.0
) -> List[Dict[str, Any]]:
    """Vyparsuje z textu přednášky značky [KAPITOLA: XX | Název] a dopočítá časové razítko."""
    pattern = r"\[KAPITOLA:\s*([0-9A-Za-z]+)\s*\|\s*([^\]]+)\]"
    matches = list(re.finditer(pattern, script))

    if not matches:
        # Fallback – vygenerujeme standardní kapitoly po úsecích
        return [
            {"time": "00:00", "seconds": 0, "title": "Klinický kontext a jádro problému"},
            {"time": "02:30", "seconds": 150, "title": "Patofyziologie v analogii"},
            {"time": "05:00", "seconds": 300, "title": "Diagnostický rozbor a tabulky"},
            {"time": "07:30", "seconds": 450, "title": "Léčebná strategie a perly"},
            {"time": "10:00", "seconds": 600, "title": "Závěrečné shrnutí"},
        ]

    total_len = len(script) if len(script) > 0 else 1
    chapters = []

    for idx, match in enumerate(matches):
        char_offset = match.start()
        ratio = char_offset / total_len
        calculated_seconds = round(ratio * total_duration_seconds, 1)

        minutes = int(calculated_seconds // 60)
        seconds = int(calculated_seconds % 60)
        time_str = f"{minutes:02d}:{seconds:02d}"

        title = match.group(2).strip()
        chapters.append(
            {
                "id": idx + 1,
                "time": time_str,
                "seconds": calculated_seconds,
                "title": title,
            }
        )

    # První kapitola vždy začíná na 00:00
    if chapters and chapters[0]["seconds"] > 0:
        chapters[0]["seconds"] = 0.0
        chapters[0]["time"] = "00:00"

    return chapters


def clean_script_for_tts(script: str) -> str:
    """Odstraní značky kapitol a metadat z textu před odesláním do TTS syntezátoru."""
    clean = re.sub(r"\[KAPITOLA:[^\]]+\]", "", script)
    clean = re.sub(r"<!--.*?-->", "", clean, flags=re.DOTALL)
    # Odstranění markdown značek pro plynulý přednes TTS
    clean = re.sub(r"\*\*([^*]+)\*\*", r"\1", clean)
    clean = re.sub(r"\*([^*]+)\*", r"\1", clean)
    clean = re.sub(r"^#+\s*", "", clean, flags=re.MULTILINE)
    # Vyčištění nadbytečných prázdných řádků
    clean = re.sub(r"\n{3,}", "\n\n", clean)
    return clean.strip()


def is_retryable_gemini_error(error: Exception) -> bool:
    """Zjistí, zda se jedná o dočasnou/řešitelnou chybu Gemini API (přetížení, rate limit, timeout)."""
    normalized = str(error).lower()
    return any(
        marker in normalized
        for marker in (
            "503",
            "unavailable",
            "overloaded",
            "429",
            "resource exhausted",
            "resource_exhausted",
            "quota",
            "timeout",
            "timed out",
            "504",
            "500",
            "internal server error",
            "connecterror",
            "connection reset",
            "remote disconnected",
            "deadline exceeded",
        )
    )


def get_fallback_model(model: str) -> Optional[str]:
    """Vrátí spolehlivý záložní model v případě trvalého přetížení vybraného modelu."""
    fallbacks = {
        "gemini-3.8-flash": "gemini-3.6-flash",
        "gemini-pro-latest": "gemini-3.6-flash",
        "gemini-flash-latest": "gemini-3.6-flash",
        "gemini-3.6-flash": "gemini-flash-latest",
    }
    return fallbacks.get(model)


async def call_gemini_with_resilience(
    ai_client: Any,
    model: str,
    contents: List[Any],
    config: types.GenerateContentConfig,
    report_fn: Optional[Any] = None,
    step_name: str = "generating_markdown",
    step_pct: int = 30,
    wait_intervals: Optional[List[int]] = None,
) -> Any:
    """
    Spolehlivé volání Gemini API s automatickými retries při přetížení (HTTP 503 / 429),
    progresivním čekáním, průběžnými notifikacemi pro uživatele a záložním modelem (fallback).
    """
    if wait_intervals is None:
        wait_intervals = [6, 12, 25, 45, 60]

    current_model = model
    fallback_used = False
    max_attempts = len(wait_intervals) + 1

    for attempt in range(max_attempts):
        try:
            resp = await asyncio.to_thread(
                ai_client.models.generate_content,
                model=current_model,
                contents=contents,
                config=config,
            )
            return resp
        except Exception as e:
            if is_retryable_gemini_error(e):
                err_detail = str(e)
                try:
                    from main import extract_gemini_error_detail
                    err_detail = extract_gemini_error_detail(e)
                except Exception:
                    pass

                # Pokud primární model selhává opakovaně (od 3. pokusu), zkusíme záložní model
                if attempt >= 2 and not fallback_used:
                    fallback = get_fallback_model(current_model)
                    if fallback and fallback != current_model:
                        fallback_used = True
                        if report_fn:
                            await report_fn(
                                step_name,
                                step_pct,
                                f"⚠️ Model {current_model} je na straně Googlu dočasně přetížen ({err_detail}). Přepínám na záložní model {fallback}...",
                            )
                        current_model = fallback
                        await asyncio.sleep(4)
                        continue

                if attempt < max_attempts - 1:
                    wait_time = wait_intervals[attempt]
                    if report_fn:
                        await report_fn(
                            step_name,
                            step_pct,
                            f"⚠️ Gemini API je dočasně přetížené ({err_detail}). Pokus {attempt + 1}/{max_attempts}, čekám {wait_time} s...",
                        )
                    await asyncio.sleep(wait_time)
                else:
                    raise
            else:
                raise


async def generate_lesson_package(
    project_id: str,
    lesson_title: str,
    target_language: str,
    file_paths: List[Tuple[str, str]],
    gemini_model: str = "gemini-3.6-flash",
    tts_provider: str = "openai",
    tts_voice: str = "onyx",
    progress_callback: Optional[Any] = None,
) -> Dict[str, Any]:
    """
    Kompletní orchestrace generování výukové lekce od A do Z:
    1. Multimodální extrakce z nahraných souborů (PDF, skeny, obrázky)
    2. Stage A: Generování strukturovaného studijního textu (Markdown)
    3. Stage B: Generování scénáře mluvené přednášky
    4. TTS syntéza mluveného audia (lecture_audio.mp3)
    5. Spočtení časových značek kapitol a uložení do SQLite
    """
    from main import (
        internal_generate_audio,
        send_log,
        AUDIO_DIR,
        NOTES_DIR,
        UPLOAD_DIR,
    )

    safe_proj = sanitize_project_name(project_id)
    ai = get_ai_client()

    async def report(step: str, pct: int, msg: str):
        await send_log(f"🎓 [Lekce '{lesson_title}'] ({pct}%): {msg}")
        if progress_callback:
            try:
                await progress_callback(
                    {
                        "step": step,
                        "percentage": pct,
                        "message": msg,
                        "project": safe_proj,
                    }
                )
            except Exception:
                pass

    # --- KROK 1: EXTRAKCE A PŘÍPRAVA MULTIMODÁLNÍCH VSTUPŮ ---
    await report("extracting", 10, "Extrahuji a analyzuji multimodální podklady (PDF, obrázky, text)...")
    multimodal_parts, extracted_text, file_metas = await extract_multimodal_parts(file_paths)

    # Uložíme extrahovaný text do ChromaDB pro budoucí chat
    if extracted_text.strip():
        await index_lesson_text_to_chroma(safe_proj, extracted_text, lesson_title)

    # --- KROK 2: STAGE A – GENERACE STRUKTUROVANÉHO STUDIJNÍHO TEXTU ---
    await report(
        "generating_markdown",
        30,
        f"Generuji komplexní studijní materiál v jazyce: {target_language} přes {gemini_model}...",
    )

    prompt_a = STAGE_A_STUDY_MATERIAL_PROMPT.replace("{TOPIC}", lesson_title).replace(
        "{TARGET_LANGUAGE}", target_language
    )

    contents_a: List[Any] = []
    # Přidáme multimodální části (obrázky, PDF)
    contents_a.extend(multimodal_parts)
    # Přidáme textový extrakt a zadání
    user_payload_a = f"TÉMA VÝUKOVÉ LEKCE: {lesson_title}\n\n"
    if extracted_text.strip():
        # Poskytneme bohatý textový kontext (až 150 000 znaků, aby nebyly oříznuty otázky na konci materiálů)
        user_payload_a += f"=== TEXTOVÝ EXTRAKT Z NAHRANÝCH SOUBORŮ ===\n{extracted_text[:150000]}\n==========================================\n\n"
    user_payload_a += (
        "Vytvoř kompletní a detailní studijní text podle zadané osnovy.\n\n"
        "DŮLEŽITÉ POKYNY:\n"
        "1. Vynech jakýkoliv zdlouhavý úvod či obecnou omáčku v úvodní sekci – jdi přímo k definici a odbornému konceptu.\n"
        "2. Důkladně prozkoumej celé přiložené podklady VČETNĚ FOTOGRAFIÍ SLIDŮ A PREZENTACÍ (např. kazuistiky Cas clinique, otázky na závěrečných snímcích).\n"
        "3. Pokud na fotografiích slidů najdeš testové otázky (MCQ s volbami A/B/C/D/E), kde jsou odpovědi zakroužkované perem/fixem, podtržené nebo doplněné rukopisem:\n"
        "   - Tyto otázky a jejich možnosti kompletně zařaď do Sekce 8 v cílovém jazyce.\n"
        "   - Zakroužkovanou volbu uveď jako správnou odpověď ze slidu a pro všechny možnosti uveď detailní zdůvodnění.\n"
        "4. Dbej na to, abys vygeneroval kompletní text všech 8 sekcí osnovy a text se nikde neusekl."
    )
    contents_a.append(user_payload_a)

    config_a = types.GenerateContentConfig(
        system_instruction=prompt_a,
        temperature=0.3,
        max_output_tokens=32768,
        thinking_config=types.ThinkingConfig(thinking_budget=2048),
    )

    resp_a = await call_gemini_with_resilience(
        ai_client=ai,
        model=gemini_model,
        contents=contents_a,
        config=config_a,
        report_fn=report,
        step_name="generating_markdown",
        step_pct=30,
    )
    study_text_markdown = resp_a.text or ""

    # Záchranná pojistka: pokud výstup dosáhl limitu tokenů, automaticky navážeme a dopíšeme zbývající sekce
    candidate = resp_a.candidates[0] if resp_a.candidates else None
    finish_reason = getattr(candidate, "finish_reason", None)
    if finish_reason and "MAX_TOKENS" in str(finish_reason):
        await send_log(f"⚠️ Výstup studijního textu dosáhl limitu tokenů ({finish_reason}). Automaticky navazuji a dopisuji zbývající sekce...")
        cont_prompt = (
            f"Předchozí výstup byl useknut kvůli limitu tokenů na tomto místě:\n"
            f"\"...{study_text_markdown[-300:]}\"\n\n"
            f"Navazuj přesně tam, kde text skončil (bez opakování předchozího textu) "
            f"a dopiš všechny zbývající sekce osnovy až do úplného konce včetně Sekce 8 s testovými otázkami a zakroužkovanými odpověďmi ze slidů."
        )
        resp_cont = await call_gemini_with_resilience(
            ai_client=ai,
            model=gemini_model,
            contents=[*contents_a, types.Part.from_text(text=study_text_markdown), cont_prompt],
            config=config_a,
            report_fn=report,
            step_name="generating_markdown",
            step_pct=45,
        )
        if resp_cont.text:
            study_text_markdown += "\n" + resp_cont.text.strip()

    if not study_text_markdown.strip():
        raise ValueError("Model nevrátil žádný studijní text.")

    # Uložení Markdownu na disk do uploads/{safe_proj}/lesson_material.md i generated_notes
    proj_upload_dir = os.path.join(UPLOAD_DIR, safe_proj)
    os.makedirs(proj_upload_dir, exist_ok=True)
    lesson_md_path = os.path.join(proj_upload_dir, "lesson_material.md")
    with open(lesson_md_path, "w", encoding="utf-8") as f:
        f.write(study_text_markdown)

    notes_backup_path = os.path.join(NOTES_DIR, f"{safe_proj}_lesson_material.md")
    with open(notes_backup_path, "w", encoding="utf-8") as f:
        f.write(study_text_markdown)

    # Zaindexujeme i výsledný studijní text do Chroma pro okamžitou přesnost chatu
    await index_lesson_text_to_chroma(safe_proj, study_text_markdown, f"{lesson_title} (Studijní text)")

    # --- KROK 3: STAGE B – GENERACE SCÉNÁŘE MLUVENÉ PŘEDNÁŠKY ---
    await report(
        "generating_lecture",
        60,
        "Vytvářím didaktický scénář přednášky profesora s kapitolami a analogiemi...",
    )

    prompt_b = STAGE_B_LECTURE_SCRIPT_PROMPT.replace("{TOPIC}", lesson_title).replace(
        "{TARGET_LANGUAGE}", target_language
    )

    user_payload_b = (
        f"TÉMA PŘEDNÁŠKY: {lesson_title}\n\n"
        f"=== HOTOVÝ STUDIJNÍ TEXT, KTERÝ MÁ STUDENT PŘED SEBOU ===\n"
        f"{study_text_markdown}\n"
        f"========================================================\n\n"
        "Vytvoř mluvený přednáškový výklad profesora s kapitolami [KAPITOLA: XX | Název].\n\n"
        "DŮLEŽITÉ POKYNY PRO PŘEDNÁŠKU:\n"
        "1. VYNECH ZDLOUHAVÝ ÚVOD: Začni maximálně jednou svižnou větou a okamžitě jdi do klinického problému a výkladu.\n"
        "2. POUŽÍVEJ SKUTEČNÝ VÝUKOVÝ PŘEDNÁŠKOVÝ STYL: Mluv živě k publiku, vysvětluj 'proč' pomocí analogií, upozorňuj na klinické chytáky.\n"
        "3. AKTIVNĚ ODKAZUJ NA TEXT: Naváděj studenta na konkrétní sekce, tabulky a schémata ve studijním materiálu.\n"
        "4. TESTOVÉ OTÁZKY: Pokud jsou na konci textu testové otázky, v závěru na ně studenta pouze stručně odkaž pro samostudium – v nahrávce je mechanicky nečti."
    )

    config_b = types.GenerateContentConfig(
        system_instruction=prompt_b,
        temperature=0.7,
        max_output_tokens=16384,
        thinking_config=types.ThinkingConfig(thinking_budget=1024),
    )

    resp_b = await call_gemini_with_resilience(
        ai_client=ai,
        model=gemini_model,
        contents=[user_payload_b],
        config=config_b,
        report_fn=report,
        step_name="generating_lecture",
        step_pct=60,
    )
    lecture_script = resp_b.text or ""

    if not lecture_script.strip():
        raise ValueError("Model nevrátil scénář přednášky.")

    # Uložíme i scénář na disk
    script_path = os.path.join(proj_upload_dir, "lecture_script.txt")
    with open(script_path, "w", encoding="utf-8") as f:
        f.write(lecture_script)

    # --- KROK 4: TTS SYNTÉZA MLUVENÉHO AUDIA ---
    await report(
        "synthesizing_audio",
        75,
        f"Spouštím syntézu mluvené přednášky přes {tts_provider.upper()} (hlas: {tts_voice})...",
    )

    clean_script = clean_script_for_tts(lecture_script)
    audio_base_filename = f"{safe_proj}_lecture_audio"

    # Voláme existující funkci syntézy zvuku v main.py
    audio_url, audio_filename, _ = await internal_generate_audio(
        script=clean_script,
        filename=audio_base_filename,
        provider=tts_provider,
        voice=tts_voice,
        output_format="mp3",
        question_title=f"Výuková lekce: {lesson_title}",
    )

    # --- KROK 5: KALKULACE ČASŮ KAPITOL A ULOŽENÍ METADAT ---
    await report("finalizing", 95, "Dopočítávám časové značky kapitol pro audio přehrávač...")

    # Zjištění přesné délky vygenerovaného audia
    full_audio_path = os.path.join(AUDIO_DIR, audio_filename)
    audio_duration = 600.0
    if os.path.exists(full_audio_path):
        try:
            mp3_info = MP3(full_audio_path)
            audio_duration = mp3_info.info.length
        except Exception:
            pass

    chapters = extract_chapters_from_script(lecture_script, total_duration_seconds=audio_duration)

    # Uložení do SQLite
    save_or_update_lesson(
        project_id=safe_proj,
        title=lesson_title,
        target_language=target_language,
        markdown_content=study_text_markdown,
        lecture_script=lecture_script,
        audio_filename=audio_filename,
        audio_url=audio_url,
        chapters=chapters,
        status="completed",
    )

    # Uložení JSON metadat i do souboru v projektu pro zálohu
    meta_payload = {
        "project_id": safe_proj,
        "title": lesson_title,
        "target_language": target_language,
        "audio_filename": audio_filename,
        "audio_url": audio_url,
        "audio_duration_seconds": audio_duration,
        "chapters": chapters,
        "sources": file_metas,
    }
    with open(os.path.join(proj_upload_dir, "lesson_metadata.json"), "w", encoding="utf-8") as f:
        json.dump(meta_payload, f, ensure_ascii=False, indent=2)

    await report("complete", 100, "Výuková lekce byla úspěšně vygenerována a připravena k poslechu!")

    return {
        "status": "completed",
        "project_id": safe_proj,
        "title": lesson_title,
        "target_language": target_language,
        "audio_url": audio_url,
        "audio_filename": audio_filename,
        "chapters": chapters,
        "markdown_content": study_text_markdown,
        "lecture_script": lecture_script,
    }
