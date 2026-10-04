"""
RAG služba pro extrakci textu ze souborů (PDF, DOCX, PPTX, TXT),
sémantický chunking, hybridní vyhledávání v ChromaDB a FTS5.
"""

import asyncio
import os
import re
from typing import Any, Tuple

try:
    import docx
except ImportError:
    docx = None

try:
    import pptx
except ImportError:
    pptx = None

try:
    import pdfplumber
except ImportError:
    pdfplumber = None

from core.clients import get_chroma_client
from core.logger import send_log
from core.utils import sanitize_name


async def extract_sections_from_file(file_path: str, filename: str) -> list[dict[str, Any]]:
    """
    Extrahuje strukturovaný text s metadaty (číslo stránky, snímku nebo sekce).
    Vrací seznam dictů: [{"text": str, "source": filename, "page": str}].
    """
    ext = filename.lower().split(".")[-1]
    sections: list[dict[str, Any]] = []

    if ext == "pdf":
        try:
            import pypdf

            reader = pypdf.PdfReader(file_path)
            for idx, page in enumerate(reader.pages):
                if idx % 25 == 0:
                    await asyncio.sleep(0.005)
                text = page.extract_text()
                if text and text.strip():
                    sections.append({"text": text.strip(), "source": filename, "page": str(idx + 1)})
        except Exception:
            if pdfplumber is not None:
                with pdfplumber.open(file_path) as pdf:
                    for idx, page in enumerate(pdf.pages):
                        if idx % 10 == 0:
                            await asyncio.sleep(0.01)
                        text = page.extract_text()
                        if text and text.strip():
                            sections.append({"text": text.strip(), "source": filename, "page": str(idx + 1)})

    elif ext == "docx":
        if docx is None:
            raise ImportError("Chybí knihovna python-docx")
        doc = docx.Document(file_path)
        current_block = []
        current_heading = "Úvod"
        for idx, para in enumerate(doc.paragraphs):
            if idx % 100 == 0:
                await asyncio.sleep(0.01)
            p_text = para.text.strip()
            if not p_text:
                continue
            if para.style and para.style.name.startswith("Heading"):
                if current_block:
                    sections.append({"text": "\n".join(current_block), "source": filename, "page": current_heading})
                    current_block = []
                current_heading = p_text[:40]
            current_block.append(p_text)
        if current_block:
            sections.append({"text": "\n".join(current_block), "source": filename, "page": current_heading})

    elif ext == "pptx":
        if pptx is None:
            raise ImportError("Chybí knihovna python-pptx")
        prs = pptx.Presentation(file_path)
        for idx, slide in enumerate(prs.slides):
            if idx % 5 == 0:
                await asyncio.sleep(0.01)
            slide_lines = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text and shape.text.strip():
                    slide_lines.append(shape.text.strip())
            if slide_lines:
                sections.append({"text": "\n".join(slide_lines), "source": filename, "page": f"Slide {idx + 1}"})

    elif ext in ("txt", "md"):
        with open(file_path, encoding="utf-8") as f:
            content = f.read()
        raw_parts = re.split(r"\n(?=#{1,3}\s+|\n---+\n)", content)
        page_counter = 1
        for part in raw_parts:
            part_clean = part.strip()
            if part_clean:
                sections.append({"text": part_clean, "source": filename, "page": str(page_counter)})
                page_counter += 1
    else:
        raise ValueError(f"Nepodporovaný formát: {ext}")

    return sections


async def extract_text_from_file(file_path: str, filename: str) -> str:
    sections = await extract_sections_from_file(file_path, filename)
    return "\n\n".join(s["text"] for s in sections)


def chunk_sections_with_metadata(
    sections: list[dict[str, Any]], chunk_size: int = 7000, overlap: int = 1000
) -> list[dict[str, Any]]:
    chunks: list[dict[str, Any]] = []
    current_text = ""
    current_source = ""
    start_page = None
    end_page = None

    def flush_chunk(text_to_save: str, src: str, p_start: Any, p_end: Any):
        clean_t = text_to_save.strip()
        if not clean_t:
            return
        page_label = str(p_start) if p_start == p_end or not p_end else f"{p_start}–{p_end}"
        chunks.append(
            {
                "document": clean_t,
                "source": src,
                "page": page_label,
            }
        )

    for sec in sections:
        sec_text = sec["text"]
        source = sec["source"]
        page = sec.get("page", "1")

        if not current_text:
            current_text = sec_text
            current_source = source
            start_page = page
            end_page = page
            continue

        if len(current_text) + len(sec_text) + 2 <= chunk_size:
            current_text += "\n\n" + sec_text
            end_page = page
        else:
            flush_chunk(current_text, current_source, start_page, end_page)

            overlap_text = ""
            if len(current_text) > overlap:
                overlap_start = len(current_text) - overlap
                split_idx = current_text.find("\n\n", overlap_start)
                if split_idx == -1:
                    split_idx = current_text.find(". ", overlap_start)
                    if split_idx != -1:
                        split_idx += 2
                if split_idx == -1:
                    split_idx = overlap_start
                overlap_text = current_text[split_idx:].strip()

            if overlap_text:
                current_text = overlap_text + "\n\n" + sec_text
                start_page = end_page
                end_page = page
            else:
                current_text = sec_text
                start_page = page
                end_page = page

    if current_text.strip():
        flush_chunk(current_text, current_source, start_page, end_page)

    return chunks


async def index_file_to_chroma(file_path: str, filename: str, project_name: str):
    await send_log(f"📖 Extraktuji text z dokumentu s analýzou stran: {filename} ({project_name})...")
    try:
        sections = await extract_sections_from_file(file_path, filename)
        if not sections:
            await send_log(f"⚠️ Dokument {filename} neobsahuje žádný strojově čitelný text.")
            return

        chunks = chunk_sections_with_metadata(sections)
        chroma = get_chroma_client()
        collection_name = f"proj_{sanitize_name(project_name)}"
        collection = chroma.get_or_create_collection(name=collection_name)

        docs = [c["document"] for c in chunks]
        ids = [f"{filename}_chunk_{i}" for i in range(len(chunks))]
        metadatas = [
            {"source": c["source"], "page": str(c.get("page", "1")), "chunk_index": i} for i, c in enumerate(chunks)
        ]

        collection.add(documents=docs, metadatas=metadatas, ids=ids)
        try:
            from chat_service import index_chunks_to_fts

            fts_chunks = [
                {
                    "id": ids[i],
                    "document": docs[i],
                    "source": metadatas[i]["source"],
                    "page": metadatas[i]["page"],
                }
                for i in range(len(chunks))
            ]
            index_chunks_to_fts(project_name, fts_chunks)
        except Exception as e:
            print(f"[RAG] FTS indexing warning: {e}")
        await send_log(
            f"✅ Dokument {filename} úspěšně zpracován ({len(chunks)} sémantických chunků s metadaty stran uloženo)."
        )
    except Exception as e:
        await send_log(f"❌ Selhalo indexování souboru {filename}: {str(e)}")


def decompose_question_to_subqueries(question: str) -> list[str]:
    """
    Rozloží kombinovanou zkouškovou otázku (např. 'Lístek 1: a) Astma bronchiale, b) Lymfomy')
    na jednotlivá dílčí témata, aby RAG našel dostatek materiálu pro všechny části zkouškového lístku.
    """
    clean_q = question.strip()
    parts = re.split(r"(?:^|\s+)(?:[a-zA-Z0-9]\)|\([a-zA-Z0-9]\))\s*", clean_q)
    sub_queries = []
    for p in parts:
        p_clean = p.strip(" :;,.-")
        if re.match(r"^(?:Lístek|Otázka|Téma|Ot\.?)\s*\d+$", p_clean, re.IGNORECASE):
            continue
        if len(p_clean) >= 3:
            sub_queries.append(p_clean)

    if not sub_queries and (";" in clean_q or "\n" in clean_q):
        for p in re.split(r"[\n;]+", clean_q):
            p_clean = p.strip(" :;,.-")
            if len(p_clean) >= 3 and not re.match(r"^(?:Lístek|Otázka)\s*\d+$", p_clean, re.IGNORECASE):
                sub_queries.append(p_clean)

    queries = [clean_q]
    for sq in sub_queries:
        if sq.lower() not in [q.lower() for q in queries]:
            queries.append(sq)
    return queries


async def query_rag_context_with_sources(
    question: str, project: str, n_results: int = 40
) -> tuple[str, list[dict[str, Any]], str]:
    safe_proj = sanitize_name(project)
    await send_log(f"🔍 Hybridní RAG: Prohledávám materiály pro téma: '{question[:50]}...'")
    from chat_service import retrieve_chat_context

    sub_queries = decompose_question_to_subqueries(question)
    query_to_use = " ".join(sub_queries) if sub_queries else question

    context_text, unique_sources, raw_context = await retrieve_chat_context(
        project_id=safe_proj,
        query=query_to_use,
        n_results=n_results,
        return_raw=True,
    )

    if not raw_context.strip():
        raise ValueError("V materiálech nebyly nalezeny žádné podklady k této otázce.")

    return context_text, unique_sources, raw_context
